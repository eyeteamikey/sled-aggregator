import csv
import hashlib
import io
import re
import stat
import zipfile
from abc import ABC, abstractmethod
from html.parser import HTMLParser
from pathlib import PurePosixPath
from xml.etree import ElementTree

from sled_aggregator.config import Settings
from sled_aggregator.documents.extraction_types import (
    ExtractedBlock,
    ExtractedTable,
    ExtractionFailure,
    ExtractionResult,
    ExtractionState,
)


class DocumentParser(ABC):
    parser_version = "1.0"
    supported_media_types: tuple[str, ...] = ()
    supported_extensions: tuple[str, ...] = ()

    @property
    def parser_name(self) -> str:
        return type(self).__name__

    @abstractmethod
    def sniff(self, data: bytes) -> bool: ...

    @abstractmethod
    def parse(self, data: bytes, *, settings: Settings) -> ExtractionResult: ...

    def result(self, detected_format: str) -> ExtractionResult:
        return ExtractionResult(self.parser_name, self.parser_version, detected_format)


class PlainTextParser(DocumentParser):
    supported_media_types = ("text/plain",)
    supported_extensions = ("txt",)

    def sniff(self, data: bytes) -> bool:
        return b"\x00" not in data

    def parse(self, data: bytes, *, settings: Settings) -> ExtractionResult:
        result = self.result("txt")
        text = data.decode("utf-8-sig", "replace")
        result.blocks = [
            ExtractedBlock(raw_text=p, paragraph_index=i)
            for i, p in enumerate(re.split(r"\n\s*\n", text))
        ]
        return result


class DelimitedTextParser(DocumentParser):
    supported_media_types = ("text/csv", "text/tab-separated-values")
    supported_extensions = ("csv", "tsv")

    def sniff(self, data: bytes) -> bool:
        sample = data[:4096].decode("utf-8-sig", "replace")
        return "\t" in sample or ("," in sample and "\n" in sample)

    def parse(self, data: bytes, *, settings: Settings) -> ExtractionResult:
        text = data.decode("utf-8-sig", "replace")
        delimiter = "\t" if text.count("\t") > text.count(",") else ","
        result = self.result("tsv" if delimiter == "\t" else "csv")
        rows = []
        for index, row in enumerate(csv.reader(io.StringIO(text), delimiter=delimiter), 1):
            if index > settings.document_spreadsheet_max_rows:
                result.warnings.append("row limit reached")
                break
            row = row[: settings.document_spreadsheet_max_columns]
            rows.append(row)
            result.blocks.append(
                ExtractedBlock(
                    raw_text="\t".join(row),
                    block_type="table_row",
                    row_start=index,
                    row_end=index,
                    column_start=1,
                    column_end=len(row),
                    table_index=0,
                )
            )
        result.tables.append(ExtractedTable(0, rows))
        return result


class _SafeHtml(HTMLParser):
    allowed = {"title", "h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "pre", "td", "th"}
    ignored = {"script", "style", "form", "nav", "iframe", "noscript"}

    def __init__(self) -> None:
        super().__init__()
        self.stack = []
        self.parts: list[tuple[str, str]] = []

    def handle_starttag(self, tag, attrs):
        self.stack.append(tag)

    def handle_endtag(self, tag):
        if self.stack:
            self.stack.pop()

    def handle_data(self, data):
        if (
            self.stack
            and not any(tag in self.ignored for tag in self.stack)
            and self.stack[-1] in self.allowed
            and data.strip()
        ):
            self.parts.append((self.stack[-1], data))


class HtmlDocumentParser(DocumentParser):
    supported_media_types = ("text/html",)
    supported_extensions = ("html", "htm")

    def sniff(self, data):
        return bool(re.match(rb"\s*(?:<!doctype\s+html|<html|<head|<body)", data, re.I))

    def parse(self, data, *, settings):
        parser = _SafeHtml()
        parser.feed(data.decode("utf-8", "replace"))
        result = self.result("html")
        result.blocks = [
            ExtractedBlock(
                raw_text=text,
                block_type="heading" if tag.startswith("h") else "paragraph",
                heading_level=int(tag[1]) if tag.startswith("h") else None,
            )
            for tag, text in parser.parts
        ]
        return result


class XmlDocumentParser(DocumentParser):
    supported_media_types = ("application/xml", "text/xml")
    supported_extensions = ("xml",)

    def sniff(self, data):
        return data.lstrip().startswith((b"<?xml", b"<"))

    def parse(self, data, *, settings):
        if b"<!DOCTYPE" in data.upper() or b"<!ENTITY" in data.upper():
            raise ExtractionFailure(ExtractionState.QUARANTINED, "DTD and entities are prohibited")
        try:
            root = ElementTree.fromstring(data)
        except ElementTree.ParseError as exc:
            raise ExtractionFailure(ExtractionState.MALFORMED, "malformed XML") from exc
        result = self.result("xml")
        for index, node in enumerate(root.iter()):
            if index >= 100_000:
                raise ExtractionFailure(ExtractionState.QUARANTINED, "XML node limit exceeded")
            text = (node.text or "").strip()
            if text:
                result.blocks.append(
                    ExtractedBlock(raw_text=text, source_location={"element": node.tag[:200]})
                )
        return result


class RtfDocumentParser(DocumentParser):
    supported_media_types = ("application/rtf", "text/rtf")
    supported_extensions = ("rtf",)

    def sniff(self, data):
        return data.startswith(b"{\\rtf")

    def parse(self, data, *, settings):
        text = data.decode("latin-1", "replace")
        text = re.sub(r"\\(?:object|pict)\b.*?}", "", text, flags=re.S)
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?|[{}]", "", text)
        result = self.result("rtf")
        result.blocks = [ExtractedBlock(raw_text=text)]
        return result


def _ooxml_kind(data: bytes) -> str | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = set(archive.namelist())
    except zipfile.BadZipFile:
        return None
    if "word/document.xml" in names:
        return "docx"
    if "xl/workbook.xml" in names:
        return "xlsx"
    return "zip"


class DocxDocumentParser(DocumentParser):
    supported_media_types = (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    )
    supported_extensions = ("docx",)

    def sniff(self, data):
        return _ooxml_kind(data) == "docx"

    def parse(self, data, *, settings):
        result = self.result("docx")
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                xml = z.read("word/document.xml")
        except (zipfile.BadZipFile, KeyError) as exc:
            raise ExtractionFailure(ExtractionState.MALFORMED, "malformed DOCX") from exc
        root = ElementTree.fromstring(xml)
        ns = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
        for i, p in enumerate(root.iter(ns + "p")):
            text = "".join(t.text or "" for t in p.iter(ns + "t"))
            if text:
                result.blocks.append(ExtractedBlock(raw_text=text, paragraph_index=i))
        return result


class SpreadsheetDocumentParser(DocumentParser):
    supported_media_types = ("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",)
    supported_extensions = ("xlsx",)

    def sniff(self, data):
        return _ooxml_kind(data) == "xlsx"

    def parse(self, data, *, settings):
        result = self.result("xlsx")
        try:
            import openpyxl

            workbook = openpyxl.load_workbook(
                io.BytesIO(data), read_only=True, data_only=False, keep_links=False
            )
        except (ImportError, OSError, ValueError, zipfile.BadZipFile) as exc:
            raise ExtractionFailure(
                ExtractionState.MALFORMED, "XLSX parser unavailable or malformed"
            ) from exc
        cells = 0
        for sheet in workbook.worksheets[: settings.document_spreadsheet_max_sheets]:
            rows = []
            for row_number, row in enumerate(
                sheet.iter_rows(
                    max_row=settings.document_spreadsheet_max_rows,
                    max_col=settings.document_spreadsheet_max_columns,
                ),
                1,
            ):
                values = ["" if c.value is None else str(c.value) for c in row]
                cells += len(values)
                if cells > settings.document_spreadsheet_max_cells:
                    result.warnings.append("cell limit reached")
                    break
                if any(values):
                    rows.append(values)
                    result.blocks.append(
                        ExtractedBlock(
                            raw_text="\t".join(values),
                            block_type="table_row",
                            sheet_name=sheet.title,
                            row_start=row_number,
                            row_end=row_number,
                            column_start=1,
                            column_end=len(values),
                            table_index=len(result.tables),
                        )
                    )
            if rows:
                result.tables.append(
                    ExtractedTable(len(result.tables), rows, sheet_name=sheet.title)
                )
        workbook.close()
        return result


class ZipDocumentParser(DocumentParser):
    supported_media_types = ("application/zip",)
    supported_extensions = ("zip",)

    def __init__(self, registry=None):
        self.registry = registry

    def sniff(self, data):
        return data.startswith(b"PK\x03\x04") and _ooxml_kind(data) == "zip"

    def parse(self, data, *, settings):
        result = self.result("zip")
        total = 0
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            infos = archive.infolist()
            if len(infos) > settings.document_zip_max_entries:
                raise ExtractionFailure(
                    ExtractionState.QUARANTINED, "ZIP entry count limit exceeded"
                )
            for info in infos:
                path = PurePosixPath(info.filename)
                mode = info.external_attr >> 16
                unsafe = (
                    path.is_absolute()
                    or ".." in path.parts
                    or re.match(r"^[A-Za-z]:", info.filename)
                    or stat.S_ISLNK(mode)
                    or bool(info.flag_bits & 1)
                )
                ratio = info.file_size / max(info.compress_size, 1)
                total += info.file_size
                if (
                    unsafe
                    or info.file_size > settings.document_zip_max_entry_bytes
                    or total > settings.document_zip_max_total_bytes
                    or ratio > settings.document_zip_max_compression_ratio
                ):
                    raise ExtractionFailure(
                        ExtractionState.QUARANTINED, "unsafe or over-limit ZIP entry"
                    )
                if info.is_dir():
                    continue
                payload = archive.read(info)
                if payload.startswith(b"PK"):
                    result.warnings.append(f"nested archive skipped: {info.filename[:200]}")
                    continue
                try:
                    selection = (
                        self.registry.select(payload, filename=info.filename)
                        if self.registry
                        else None
                    )
                except ExtractionFailure:
                    selection = None
                if not selection:
                    result.warnings.append(f"unsupported archive entry: {info.filename[:200]}")
                    continue
                child = selection.parser.parse(payload, settings=settings)
                for block in child.blocks:
                    block.archive_entry = info.filename
                result.blocks.extend(child.blocks)
                result.tables.extend(child.tables)
                result.metadata.setdefault("entries", []).append(
                    {
                        "path": info.filename[:500],
                        "sha256": hashlib.sha256(payload).hexdigest(),
                        "format": child.detected_format,
                    }
                )
        return result


class PdfDocumentParser(DocumentParser):
    supported_media_types = ("application/pdf",)
    supported_extensions = ("pdf",)

    def sniff(self, data):
        return data.startswith(b"%PDF-")

    def parse(self, data, *, settings):
        try:
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data), strict=True)
            if reader.is_encrypted:
                raise ExtractionFailure(ExtractionState.UNSUPPORTED, "encrypted PDF unsupported")
        except ImportError as exc:
            raise ExtractionFailure(ExtractionState.UNSUPPORTED, "pypdf is not installed") from exc
        except ExtractionFailure:
            raise
        except Exception as exc:
            raise ExtractionFailure(ExtractionState.MALFORMED, "malformed PDF") from exc
        if len(reader.pages) > settings.document_pdf_max_pages:
            raise ExtractionFailure(ExtractionState.QUARANTINED, "PDF page limit exceeded")
        result = self.result("pdf")
        decisions = []
        for number, page in enumerate(reader.pages, 1):
            try:
                text = page.extract_text() or ""
            except Exception:
                text = ""
                result.warnings.append(f"native extraction failed on page {number}")
            images = len(page.images) > 0
            meaningful = sum(c.isalnum() for c in text)
            words = len(text.split())
            ratio = text.count("\ufffd") / max(len(text), 1)
            usable = (
                meaningful >= settings.document_pdf_native_text_min_characters
                and words >= settings.document_pdf_native_text_min_words
                and ratio <= settings.document_pdf_max_replacement_ratio
            )
            decision = "native_text" if usable else ("OCR_required" if images else "blank")
            decisions.append(
                {
                    "page": number,
                    "decision": decision,
                    "meaningful_characters": meaningful,
                    "words": words,
                    "replacement_ratio": ratio,
                }
            )
            if text.strip():
                result.blocks.append(ExtractedBlock(raw_text=text, page_number=number))
        required = sum(d["decision"] == "OCR_required" for d in decisions)
        if required:
            result.ocr_state = result.ocr_state.UNAVAILABLE
            result.warnings.append("page-specific OCR required; renderer/provider unavailable")
        result.metadata["page_decisions"] = decisions
        return result


class ParserSelection:
    def __init__(self, parser, warnings=()):
        self.parser, self.warnings = parser, list(warnings)


class DocumentParserRegistry:
    def __init__(self):
        self.parsers = [
            PdfDocumentParser(),
            DocxDocumentParser(),
            SpreadsheetDocumentParser(),
            HtmlDocumentParser(),
            XmlDocumentParser(),
            RtfDocumentParser(),
            DelimitedTextParser(),
            PlainTextParser(),
        ]
        archive = ZipDocumentParser(self)
        self.parsers.insert(3, archive)

    def select(
        self, data: bytes, *, detected_media_type=None, declared_media_type=None, filename=None
    ):
        if not data:
            raise ExtractionFailure(ExtractionState.MALFORMED, "empty artifact")
        confident = next((p for p in self.parsers if p.sniff(data)), None)
        parser = confident or next(
            (
                p
                for p in self.parsers
                if detected_media_type in p.supported_media_types and p.sniff(data)
            ),
            None,
        )
        if not parser:
            raise ExtractionFailure(ExtractionState.UNSUPPORTED, "unsupported document format")
        warnings = []
        extension = PurePosixPath(filename or "").suffix.lower().lstrip(".")
        if extension and extension not in parser.supported_extensions:
            warnings.append(
                f"extension mismatch: .{extension} content selected {parser.parser_name}"
            )
        if (
            detected_media_type
            and detected_media_type not in parser.supported_media_types
            and parser is not PlainTextParser
        ):
            warnings.append(f"media type mismatch: {detected_media_type}")
        return ParserSelection(parser, warnings)
